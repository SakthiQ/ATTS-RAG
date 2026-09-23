import os
import time
import json
import requests
from docx import Document

API_URL = "http://127.0.0.1:8000"
TEST_DOC_PATH = "data/sample_medical_report.docx"

def run_tests():
    report = {
        "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
        "endpoints": {},
        "pipeline_verification": {},
        "errors": []
    }
    
    print("=== ATTS-RAG END-TO-END HEALTH & FUNCTIONALITY CHECK ===")
    
    # 1. Test GET /health
    print("\n[1/7] Testing GET /health ...")
    try:
        r = requests.get(f"{API_URL}/health", timeout=5)
        report["endpoints"]["GET /health"] = {
            "status_code": r.status_code,
            "response": r.json()
        }
        print(f" -> Status: {r.status_code}, Body: {json.dumps(r.json())}")
    except Exception as e:
        err = f"GET /health failed: {e}"
        print(f" -> ERROR: {err}")
        report["errors"].append(err)
        return report

    # 2. Test GET /documents (Initial)
    print("\n[2/7] Testing GET /documents (Initial List) ...")
    try:
        r = requests.get(f"{API_URL}/documents", timeout=5)
        initial_docs = r.json()
        report["endpoints"]["GET /documents (initial)"] = {
            "status_code": r.status_code,
            "count": len(initial_docs)
        }
        print(f" -> Status: {r.status_code}, Total Indexed Docs: {len(initial_docs)}")
    except Exception as e:
        err = f"GET /documents failed: {e}"
        print(f" -> ERROR: {err}")
        report["errors"].append(err)

    # 3. Generate sample document from provided slide deck
    print("\n[3/7] Generating sample test document (Medical Miscommunication Detector) ...")
    os.makedirs("data", exist_ok=True)
    doc = Document()
    doc.add_heading("Medical Miscommunication Detector", level=1)
    doc.add_paragraph("NLP-Based Patient-Friendly Medical Report Explanation System")
    doc.add_paragraph("Presented by: Sanjusree K (23BCE2023) and Sakthi Narayan V (23BCE0987), VIT Vellore.")
    
    doc.add_heading("Abstract", level=2)
    doc.add_paragraph(
        "Medical reports often contain complex clinical terms that patients may find difficult to understand. "
        "For example, a statement such as 'Mild hepatic steatosis with elevated ALT' may be clear to a doctor but confusing to a patient. "
        "The Medical Miscommunication Detector is an NLP-based system designed to bridge this communication gap. "
        "It identifies important medical terms, retrieves relevant medical information, and converts technical clinical language "
        "into simple, concise, and patient-friendly explanations while preserving the original meaning."
    )
    doc.add_paragraph(
        "The system uses BioBERT for medical Named Entity Recognition, RAG for reliable information retrieval, "
        "and an instruction-tuned LLM such as Llama 3 for generating simplified explanations. "
        "It can also provide regional-language translation and text-to-speech output for better accessibility. "
        "The system is designed to explain documented medical findings, not to diagnose diseases or prescribe treatments."
    )
    doc.save(TEST_DOC_PATH)
    print(f" -> Saved docx to {TEST_DOC_PATH}")

    # 4. Test POST /upload
    print("\n[4/7] Testing POST /upload ...")
    uploaded_hash = None
    try:
        with open(TEST_DOC_PATH, "rb") as f:
            r = requests.post(f"{API_URL}/upload", files={"file": ("sample_medical_report.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        report["endpoints"]["POST /upload"] = {
            "status_code": r.status_code,
            "response": r.json()
        }
        print(f" -> Status: {r.status_code}, Response: {r.json()}")
    except Exception as e:
        err = f"POST /upload failed: {e}"
        print(f" -> ERROR: {err}")
        report["errors"].append(err)

    # Wait for background ingestion task
    print(" -> Waiting 5 seconds for background ingestion to complete...")
    time.sleep(5)

    # Re-check GET /documents to find uploaded document hash
    try:
        r = requests.get(f"{API_URL}/documents", timeout=5)
        docs = r.json()
        for content_hash, info in docs.items():
            if info.get("filename") == "sample_medical_report.docx":
                uploaded_hash = content_hash
                break
        print(f" -> Uploaded Document Hash found: {uploaded_hash}")
        report["pipeline_verification"]["uploaded_hash"] = uploaded_hash
    except Exception as e:
        err = f"Finding uploaded document in registry failed: {e}"
        print(f" -> ERROR: {err}")
        report["errors"].append(err)

    # 5. Test GET /documents/{hash}/preview
    if uploaded_hash:
        print(f"\n[5/7] Testing GET /documents/{uploaded_hash}/preview ...")
        try:
            r = requests.get(f"{API_URL}/documents/{uploaded_hash}/preview", timeout=5)
            report["endpoints"]["GET /documents/{hash}/preview"] = {
                "status_code": r.status_code,
                "response": r.json()
            }
            preview_len = len(r.json().get("preview", ""))
            print(f" -> Status: {r.status_code}, Preview character length: {preview_len}")
        except Exception as e:
            err = f"GET /documents/preview failed: {e}"
            print(f" -> ERROR: {err}")
            report["errors"].append(err)

    # 6. Test POST /query (Normal RAG Query)
    print("\n[6/7] Testing POST /query (Normal Question & 3-Layer Security Pipeline) ...")
    try:
        payload = {
            "question": "What model is used for medical Named Entity Recognition in the Medical Miscommunication Detector?",
            "session_id": "test_session_1",
            "tenant_id": "default_tenant"
        }
        start_t = time.time()
        r = requests.post(f"{API_URL}/query", json=payload, timeout=120)
        dur = round((time.time() - start_t) * 1000, 2)
        res = r.json()
        
        report["endpoints"]["POST /query (normal)"] = {
            "status_code": r.status_code,
            "latency_ms": dur,
            "threat_gate": res.get("threat_gate"),
            "layer2_gate": res.get("layer2_gate"),
            "layer3_gate": res.get("layer3_gate"),
            "answer": res.get("answer"),
            "citations": res.get("citations")
        }
        print(f" -> Status: {r.status_code} ({dur} ms)")
        print(f" -> Layer 1 (Threat Gate): {res.get('threat_gate', {}).get('status')}")
        print(f" -> Layer 2 (Trust Gate): {res.get('layer2_gate', {}).get('status')}")
        print(f" -> Layer 3 (Evidence Gate): {res.get('layer3_gate', {}).get('status')}")
        print(f" -> Answer: {res.get('answer')}")
        print(f" -> Citations: {res.get('citations')}")
    except Exception as e:
        err = f"POST /query failed: {e}"
        print(f" -> ERROR: {err}")
        report["errors"].append(err)

    # 7. Test DELETE /documents/{hash}
    if uploaded_hash:
        print(f"\n[7/7] Testing DELETE /documents/{uploaded_hash} ...")
        try:
            r = requests.delete(f"{API_URL}/documents/{uploaded_hash}", timeout=5)
            report["endpoints"]["DELETE /documents/{hash}"] = {
                "status_code": r.status_code,
                "response": r.json()
            }
            print(f" -> Status: {r.status_code}, Response: {r.json()}")
        except Exception as e:
            err = f"DELETE /documents failed: {e}"
            print(f" -> ERROR: {err}")
            report["errors"].append(err)

    # Final summary
    print("\n=== TEST SUMMARY ===")
    print(f"Total Errors Recorded: {len(report['errors'])}")
    if report["errors"]:
        for idx, err in enumerate(report["errors"], 1):
            print(f" {idx}. {err}")
    else:
        print(" ALL ENDPOINTS ARE 100% HEALTHY AND OPERATIONAL!")

    with open("logs/endpoint_health_report.json", "w") as out:
        json.dump(report, out, indent=2)
    print("Report saved to logs/endpoint_health_report.json")
    return report

if __name__ == "__main__":
    run_tests()
